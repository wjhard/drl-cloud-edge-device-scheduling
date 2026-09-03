"""Generate the final report DOCX files from docs/技术报告.md without Word COM.

This portable fallback keeps the two public Word deliverables byte-identical and
uses Word-native headings, tables, captions, inline figures, page numbering and
a table-of-contents field.  Microsoft Word can refresh the TOC on first open.
"""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "docs" / "技术报告.md"
DEFAULT_OUTPUT = ROOT / "docs" / "技术报告.docx"
DEFAULT_PUBLIC_OUTPUT = (
    ROOT
    / "docs"
    / "操作系统创新小分队_基于深度强化学习的云—边—端异构计算资源管理调度方法_项目说明书.docx"
)


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_border(parent, edge: str, value: str, size: str = "8") -> None:
    borders = parent.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        parent.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), value)
    if value == "single":
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), "000000")


def set_cell_bottom_border(cell, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:color"), "000000")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开文档后按 Ctrl+A、F9 更新目录"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, placeholder, fld_end])


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return text.strip()


def split_table_row(line: str) -> list[str]:
    return [clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(4)

    heading_settings = {
        "Heading 1": ("黑体", 16, True),
        "Heading 2": ("黑体", 14, True),
        "Heading 3": ("黑体", 12, True),
    }
    for style_name, (font_name, size, bold) in heading_settings.items():
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code" not in styles:
        style = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Code"]
    style.font.name = "Consolas"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    style.font.size = Pt(8.5)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.left_indent = Cm(0.4)
    style.paragraph_format.right_indent = Cm(0.4)
    style.paragraph_format.space_after = Pt(1)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.vertical_alignment = 1
    blocks = [
        ("第三届中国研究生操作系统开源创新大赛", 18, True, 12),
        ("暨开放原子大赛操作系统专项赛", 15, True, 24),
        ("第 16 题：云—边—端异构计算资源调度", 14, True, 40),
        ("基于深度强化学习的\n云—边—端异构计算资源管理调度方法", 23, True, 28),
        ("项 目 说 明 书", 20, True, 40),
        ("参赛队伍：操作系统创新小分队", 12, False, 8),
        ("二〇二六年八月", 12, False, 0),
    ]
    for text, size, bold, after in blocks:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(after)
        run = paragraph.add_run(text)
        set_east_asia_font(run, "黑体" if bold else "宋体")
        run.font.size = Pt(size)
        run.font.bold = bold

    doc.add_page_break()
    section.vertical_alignment = 0
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.first_line_indent = Cm(0)
    run = toc_title.add_run("目  录")
    set_east_asia_font(run, "黑体")
    run.font.size = Pt(18)
    run.font.bold = True
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_toc(toc)
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_pr = table._tbl.tblPr
    for edge, value, size in (
        ("top", "single", "12"),
        ("bottom", "single", "12"),
        ("left", "nil", "0"),
        ("right", "nil", "0"),
        ("insideH", "nil", "0"),
        ("insideV", "nil", "0"),
    ):
        set_border(table_pr, edge, value, size)
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row[column_index] if column_index < len(row) else ""
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    set_east_asia_font(run, "黑体" if row_index == 0 else "宋体")
                    run.font.size = Pt(8.5)
                    run.font.bold = row_index == 0
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                set_cell_bottom_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_equation(doc: Document, expression: str, number: int) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(2.0), Cm(12.0), Cm(2.0))
    for index, width in enumerate(widths):
        table.cell(0, index).width = width
    table_pr = table._tbl.tblPr
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        set_border(table_pr, edge, "nil", "0")

    center = table.cell(0, 1)
    center_paragraph = center.paragraphs[0]
    center_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    center_paragraph.paragraph_format.first_line_indent = Cm(0)
    omath = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = expression
    math_run.append(math_text)
    omath.append(math_run)
    center_paragraph._p.append(omath)

    right = table.cell(0, 2).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.first_line_indent = Cm(0)
    right.add_run(f"（{number}）")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, alt: str, relative_path: str, number: int) -> None:
    path = ROOT / relative_path
    if not path.is_file():
        paragraph = doc.add_paragraph(f"[图像缺失：{relative_path}]")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.add_run().add_picture(str(path), width=Cm(15.5))
    caption_text = re.sub(r"^图\s*\d+\s*", "", alt).strip()
    caption = doc.add_paragraph(f"图 {number}  {caption_text}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_after = Pt(6)


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(clean_inline(text))
    set_east_asia_font(run, "宋体")
    run.font.size = Pt(11)


def parse_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next((i for i, line in enumerate(lines) if line.strip() == "## 摘要"), 0)
    lines = lines[start:]
    index = 0
    paragraph_buffer: list[str] = []
    figure_number = 0
    equation_number = 0

    def flush() -> None:
        if paragraph_buffer:
            add_body_paragraph(doc, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer.clear()

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped.startswith("```"):
            flush()
            language = stripped[3:].strip()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if language.lower() == "math":
                equation_number += 1
                add_equation(doc, " ".join(item.strip() for item in code), equation_number)
            else:
                for code_line in code or [""]:
                    doc.add_paragraph(code_line, style="Code")
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", stripped)
        if image_match:
            flush()
            figure_number += 1
            add_figure(doc, image_match.group(1), image_match.group(2), figure_number)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            flush()
            rows = [split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush()
            raw_level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            level = 1 if text == "摘要" else min(raw_level, 3)
            doc.add_heading(text, level=level)
            index += 1
            continue
        list_match = re.match(r"^\s*(\d+\.|[-+*])\s+(.+)$", raw)
        if list_match:
            flush()
            style = "List Number" if list_match.group(1).endswith(".") else "List Bullet"
            paragraph = doc.add_paragraph(clean_inline(list_match.group(2)), style=style)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            index += 1
            continue
        if stripped.startswith(">"):
            flush()
            paragraph = doc.add_paragraph(clean_inline(stripped.lstrip("> ")))
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            for run in paragraph.runs:
                run.font.italic = True
                set_east_asia_font(run, "楷体")
            index += 1
            continue
        if stripped == "---":
            flush()
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1
    flush()


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "基于深度强化学习的云—边—端异构计算资源管理调度方法——项目说明书"
    props.subject = "第16题：云—边—端异构计算资源调度"
    props.author = "操作系统创新小分队"
    props.last_modified_by = "操作系统创新小分队"
    props.keywords = "深度强化学习, 异构计算, 云边端, 调度, openEuler"


def build(input_path: Path, output_path: Path, public_output: Path) -> None:
    markdown = input_path.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    parse_markdown(doc, markdown)
    set_core_properties(doc)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.first_line_indent = Cm(0)
        if not header.text:
            run = header.add_run("基于深度强化学习的云—边—端异构计算资源管理调度方法")
            set_east_asia_font(run, "宋体")
            run.font.size = Pt(8.5)
        footer = section.footer.paragraphs[0]
        add_page_number(footer)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    shutil.copyfile(output_path, public_output)
    print(f"generated: {output_path} ({output_path.stat().st_size} bytes)")
    print(f"copied:    {public_output} ({public_output.stat().st_size} bytes)")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--public-output", type=Path, default=DEFAULT_PUBLIC_OUTPUT)
    args = parser.parse_args()
    build(args.input, args.output, args.public_output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
